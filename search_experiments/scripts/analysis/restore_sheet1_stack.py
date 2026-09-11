#!/usr/bin/env python3
"""
One-off repair #4: Sheet1's columns A:E are a dead VSTACK array formula
(`=_xlfn.VSTACK(#REF!,#REF!)` at A1, ref="A1:E2362") whose two source-sheet
references are broken -- Excel wipes the sheet name from a broken reference,
so the formula text itself can't say what it used to point to. Every cell
in the spill range reads #REF! and the dependent chart series ("r=10",
"r=1", "Standard", "BLO") cache all zeros.

Reconstructed the original two sources from the chart's own cell ranges and
matching row counts / run_id prefixes:
- Sheet1 rows 2-1961 (1960 rows) exactly match
  'zrecord_GTR_FO_fast_findopt_reg' (zrecord_GTR_FO_fast_findopt_regular_mean)
  in full, including its own internal r10 (rows 2-879, 878 rows) / r1
  (rows 880-1961, 1082 rows) split -- which is exactly the "r=10" / "r=1"
  chart series' cell ranges.
- Sheet1 row 1962 is that second source's own repeated header row (VSTACK
  applied to whole A1:E-ref ranges, so each source's header row spills
  too), which is why row 1963 (not 1962) is where "BLO" starts.
- Sheet1 rows 1963-2362 (400 rows) exactly match
  'zrecord_GTR_FO_fast_findopt_fin' (zrecord_GTR_FO_fast_findopt_findopt_mean)
  in full; its first 200 rows (1963-2162) are exactly "BLO"'s range. The
  remaining 200 rows (2163-2362, the sheet's own r1 portion) aren't used by
  any current series.

Per the user's own fallback suggestion, this writes the literal values from
those two source sheets directly into Sheet1 (replacing the dead array
formula), rather than trying to hand-author a new array formula via raw XML
-- then also rewrites the numCache in chart2.xml ("r=10"/"r=1") and
chart3.xml ("Standard"/"BLO") so the chart displays correctly without
requiring Excel to recalculate/refresh first.

Only touches xl/worksheets/sheet6.xml (Sheet1), xl/charts/chart2.xml, and
xl/charts/chart3.xml.
"""

import re
import zipfile
import openpyxl

TARGET = "PhylogenyTreeLengths.xlsx"
HEADER = ["run_id", "candidates", "time_elapsed", "logL", "true_minus_current"]
COLS = ["A", "B", "C", "D", "E"]


def esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def fmt_num(v):
    if isinstance(v, int):
        return str(v)
    return repr(float(v))


def build_row_xml(row_num, values):
    cells = []
    for col, val in zip(COLS, values):
        ref = f"{col}{row_num}"
        if isinstance(val, str):
            cells.append(f'<c r="{ref}" t="inlineStr"><is><t>{esc(val)}</t></is></c>')
        else:
            cells.append(f'<c r="{ref}"><v>{fmt_num(val)}</v></c>')
    return f'<row r="{row_num}" spans="1:5" x14ac:dyDescent="0.3">{"".join(cells)}</row>'


def load_source(wb, sheet_name):
    ws = wb[sheet_name]
    rows = []
    for r in range(2, ws.max_row + 1):
        rows.append([ws.cell(row=r, column=c).value for c in range(1, 6)])
    return rows


def build_numcache(values):
    pts = "".join(f'<c:pt idx="{i}"><c:v>{fmt_num(v)}</c:v></c:pt>' for i, v in enumerate(values))
    return f'<c:formatCode>General</c:formatCode><c:ptCount val="{len(values)}"/>{pts}'


def replace_series_cache(xml, title, x_values, y_values):
    start = xml.find(f"<c:tx><c:v>{title}</c:v></c:tx>")
    assert start != -1, f"series {title!r} not found"
    block_start = xml.rfind("<c:ser>", 0, start)
    block_end = xml.find("</c:ser>", start) + len("</c:ser>")
    block = xml[block_start:block_end]

    new_block = re.sub(
        r"(<c:xVal><c:numRef><c:f>[^<]*</c:f><c:numCache>).*?(</c:numCache>)",
        lambda m: m.group(1) + build_numcache(x_values) + m.group(2),
        block, count=1, flags=re.DOTALL,
    )
    new_block = re.sub(
        r"(<c:yVal><c:numRef><c:f>[^<]*</c:f><c:numCache>).*?(</c:numCache>)",
        lambda m: m.group(1) + build_numcache(y_values) + m.group(2),
        new_block, count=1, flags=re.DOTALL,
    )
    assert new_block != block, f"no cache substitution happened for {title!r}"
    return xml[:block_start] + new_block + xml[block_end:]


def main():
    wb = openpyxl.load_workbook(TARGET, data_only=True)
    reg_rows = load_source(wb, "zrecord_GTR_FO_fast_findopt_reg")   # 1960 rows
    fin_rows = load_source(wb, "zrecord_GTR_FO_fast_findopt_fin")   # 400 rows
    assert len(reg_rows) == 1960 and len(fin_rows) == 400

    all_rows = [HEADER] + reg_rows + [HEADER] + fin_rows
    assert len(all_rows) == 2362

    with zipfile.ZipFile(TARGET) as z:
        entries = {n: z.read(n) for n in z.namelist()}
        order = z.namelist()

    sheet_xml = entries["xl/worksheets/sheet6.xml"].decode("utf-8")
    old_start = sheet_xml.find('<row r="1" ')
    old_end = sheet_xml.find("</row>", sheet_xml.find('<row r="2362" ')) + len("</row>")
    assert old_start != -1 and old_end != -1

    new_rows_xml = "".join(build_row_xml(i + 1, row) for i, row in enumerate(all_rows))
    sheet_xml = sheet_xml[:old_start] + new_rows_xml + sheet_xml[old_end:]
    entries["xl/worksheets/sheet6.xml"] = sheet_xml.encode("utf-8")

    # r=10 / r=1 -> reg_rows split at its internal r10/r1 boundary (row 880 in source = index 878)
    r10_reg = reg_rows[:878]
    r1_reg = reg_rows[878:]
    # BLO -> first 200 rows of fin_rows (its own r10 portion)
    blo_rows = fin_rows[:200]

    c2 = entries["xl/charts/chart2.xml"].decode("utf-8")
    c2 = replace_series_cache(c2, "r=10", [r[2] for r in r10_reg], [r[3] for r in r10_reg])
    c2 = replace_series_cache(c2, "r=1", [r[2] for r in r1_reg], [r[3] for r in r1_reg])
    entries["xl/charts/chart2.xml"] = c2.encode("utf-8")

    c3 = entries["xl/charts/chart3.xml"].decode("utf-8")
    c3 = replace_series_cache(c3, "Standard", [r[2] for r in r10_reg], [r[3] for r in r10_reg])
    c3 = replace_series_cache(c3, "BLO", [r[2] for r in blo_rows], [r[3] for r in blo_rows])
    entries["xl/charts/chart3.xml"] = c3.encode("utf-8")

    with zipfile.ZipFile(TARGET, "w", zipfile.ZIP_DEFLATED) as z:
        for name in order:
            z.writestr(name, entries[name])

    print(f"Sheet1: wrote {len(all_rows)} rows (header + {len(reg_rows)} reg + header + "
          f"{len(fin_rows)} fin), replacing the dead VSTACK/#REF! formula.")
    print("chart2.xml: refreshed numCache for r=10 (878 pts) / r=1 (1082 pts).")
    print("chart3.xml: refreshed numCache for Standard (878 pts) / BLO (200 pts).")


if __name__ == "__main__":
    main()
